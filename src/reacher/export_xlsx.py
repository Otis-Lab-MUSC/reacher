"""Excel workbook for a session export: an Events sheet and a Summary sheet.

The workbook sits alongside behavior_events*.csv in the export ZIP (the CSVs stay for
analysis scripts). The Summary sheet is computed from the very rows written to the Events
sheet, so the two cannot disagree.
"""

from __future__ import annotations

import csv
import io
from collections import Counter
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EVENT_COLUMNS = ["device", "event", "start_timestamp", "end_timestamp", "start_frame_index", "end_frame_index"]
_NUMERIC = {"start_timestamp", "end_timestamp", "start_frame_index", "end_frame_index"}
_HEADER_FILL = PatternFill("solid", fgColor="DDE6F0")


def _cell(column: str, value: Any) -> Any:
    """CSV text -> a typed cell: numeric columns become numbers, blanks stay empty."""
    if value in (None, ""):
        return None
    if column in _NUMERIC:
        try:
            return int(value)
        except (TypeError, ValueError):
            try:
                return float(value)
            except (TypeError, ValueError):
                return value
    return value


def rows_from_csv(csv_text: str) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(csv_text)))


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL


def build_behavior_workbook(
    segments: Iterable[list[dict[str, Any]]],
    facts: Optional[list[tuple[str, Any]]] = None,
) -> bytes:
    """Return .xlsx bytes.

    segments: behavior rows per segment, in order (one list for an unsegmented session).
    facts: ordered (label, value) pairs shown at the top of the Summary sheet.
    """
    segments = [list(seg) for seg in segments]
    segmented = len(segments) > 1

    wb = Workbook()
    ws = wb.active
    assert ws is not None  # a new Workbook always has an active sheet
    ws.title = "Events"
    columns = (["segment"] if segmented else []) + EVENT_COLUMNS
    ws.append(columns)
    _style_header(ws, len(columns))

    counts: Counter[str] = Counter()
    for seg_no, rows in enumerate(segments, start=1):
        for row in rows:
            counts[f"{row.get('device', '')}.{row.get('event', '')}"] += 1
            values = [_cell(c, row.get(c)) for c in EVENT_COLUMNS]
            ws.append(([seg_no] if segmented else []) + values)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(name) + 4)

    summary = wb.create_sheet("Summary")
    summary.append(["Session", ""])
    _style_header(summary, 2)
    for label, value in facts or []:
        summary.append([label, value if value not in ("",) else None])
    summary.append([])
    header_row = summary.max_row + 1
    summary.append(["Event (DEVICE.EVENT)", "Count"])
    for c in (1, 2):
        cell = summary.cell(row=header_row, column=c)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for key in sorted(counts):
        summary.append([key, counts[key]])
    total_row = summary.max_row + 1
    summary.append(["Total events", sum(counts.values())])
    for c in (1, 2):
        summary.cell(row=total_row, column=c).font = Font(bold=True)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 22
    for r in summary.iter_rows(min_row=1, max_row=summary.max_row, min_col=2, max_col=2):
        r[0].alignment = Alignment(horizontal="left")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
