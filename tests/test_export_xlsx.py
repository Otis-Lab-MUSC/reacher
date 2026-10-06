"""behavior_events.xlsx: Events + Summary sheets for the session export ZIP."""

import io

from openpyxl import load_workbook

from reacher.export_xlsx import build_behavior_workbook, rows_from_csv

_ROWS = [
    {"device": "LEVER_RH", "event": "ACTIVE", "start_timestamp": "1000", "end_timestamp": "1200", "start_frame_index": "", "end_frame_index": ""},
    {"device": "PUMP", "event": "INFUSION", "start_timestamp": "1005", "end_timestamp": "", "start_frame_index": "3", "end_frame_index": ""},
    {"device": "LEVER_RH", "event": "ACTIVE", "start_timestamp": "2000", "end_timestamp": "2100", "start_frame_index": "", "end_frame_index": ""},
]


def _open(data: bytes):
    return load_workbook(io.BytesIO(data))


def test_events_and_summary_sheets():
    wb = _open(build_behavior_workbook([_ROWS], [("Paradigm", "vi"), ("Infusions", 1)]))
    assert wb.sheetnames == ["Events", "Summary"]

    events = list(wb["Events"].iter_rows(values_only=True))
    assert events[0] == ("device", "event", "start_timestamp", "end_timestamp", "start_frame_index", "end_frame_index")
    assert events[1] == ("LEVER_RH", "ACTIVE", 1000, 1200, None, None)  # numeric cells, blanks empty
    assert len(events) == 4

    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Paradigm"] == "vi"
    assert summary["Infusions"] == 1
    assert summary["LEVER_RH.ACTIVE"] == 2
    assert summary["PUMP.INFUSION"] == 1
    assert summary["Total events"] == 3


def test_segmented_session_adds_segment_column_and_sums_all_segments():
    wb = _open(build_behavior_workbook([_ROWS[:2], _ROWS[2:]]))
    events = list(wb["Events"].iter_rows(values_only=True))
    assert events[0][0] == "segment"
    assert [r[0] for r in events[1:]] == [1, 1, 2]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Total events"] == 3


def test_empty_session_still_produces_both_sheets():
    wb = _open(build_behavior_workbook([[]]))
    assert wb.sheetnames == ["Events", "Summary"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Total events"] == 0


def test_rows_from_csv_roundtrip():
    csv_text = "device,event,start_timestamp,end_timestamp,start_frame_index,end_frame_index\nPUMP,INFUSION,5,,,\n"
    assert rows_from_csv(csv_text)[0]["event"] == "INFUSION"
